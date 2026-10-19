"""
Guided setup: from your existing resumes to a working profile, with nothing entering it unconfirmed
==================================================================================================

1. Read the resumes you upload (.docx, .pdf, .txt).
2. The local model PROPOSES roles, bullets, skills, education and certificates from each resume.
3. Python checks every proposed bullet against the resume text it came from:
     verified  - found word for word (or almost) in your resume
     reworded  - not found: the model changed it; it starts unticked and is marked for your attention
   Dates that do not appear in the resume are flagged the same way.
4. You confirm, edit or delete each item on the review page. Only what you keep is written.
5. The profile pack is built and validated: fact_bank.json, resume_profile.json, domain.toml (from a starter
   pack), search.toml and resume_template.docx.

Work in progress is saved in profiles/<name>/_setup.json, so you can stop and continue later.
"""
from __future__ import annotations

import json
import os
import re
import shutil
from datetime import datetime
from pathlib import Path
from typing import Callable, Optional

from modules.settings import PROFILES_DIR

PACKS_DIR = PROFILES_DIR / "_packs"
EXAMPLE = PROFILES_DIR / "example"
LEVELS = {"expert": "", "working": " (working knowledge)", "familiar": " (familiar)", "novice": " (novice)"}
SECTIONS = {"tools": "Tools & Software", "methods": "Methods", "domain": "Domain"}
MAX_RESUME_CHARS = 14000


# --------------------------------------------------------------------------- #
# State
# --------------------------------------------------------------------------- #

def slugify(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", (s or "").lower()).strip("_")[:40]


def folder(slug: str) -> Path:
    return PROFILES_DIR / slug


def state_path(slug: str) -> Path:
    return folder(slug) / "_setup.json"


def load_state(slug: str) -> dict:
    p = state_path(slug)
    try:
        return json.loads(p.read_text(encoding="utf-8")) if p.exists() else {}
    except (ValueError, OSError):
        return {}


def save_state(slug: str, st: dict) -> None:
    p = state_path(slug)
    p.parent.mkdir(parents=True, exist_ok=True)
    st["updated"] = datetime.now().isoformat(timespec="seconds")
    tmp = p.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(st, indent=2, ensure_ascii=False), encoding="utf-8")
    os.replace(tmp, p)


def in_progress() -> list[dict]:
    out = []
    for p in PROFILES_DIR.glob("*/_setup.json"):
        st = load_state(p.parent.name)
        out.append({"slug": p.parent.name, "name": st.get("basics", {}).get("name", p.parent.name),
                    "step": st.get("step", 1), "updated": st.get("updated", "")})
    return sorted(out, key=lambda x: x["updated"], reverse=True)


# --------------------------------------------------------------------------- #
# Reading resumes
# --------------------------------------------------------------------------- #

def read_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        import docx
        d = docx.Document(str(path))
        lines = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                lines += [c.text for c in row.cells]
        text = "\n".join(lines)
    elif suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            raise RuntimeError("Reading PDF resumes needs one free package: pip install pypdf")
        text = "\n".join((pg.extract_text() or "") for pg in PdfReader(str(path)).pages)
    else:
        text = path.read_text(encoding="utf-8", errors="ignore")
    return re.sub(r"[ \t\u00a0]+", " ", text).strip()


# --------------------------------------------------------------------------- #
# Model proposal
# --------------------------------------------------------------------------- #

EXTRACT_SYSTEM = """You turn ONE resume into structured data. You copy, you never write.
Output ONLY one JSON object, no markdown, no commentary:
{"roles":[{"employer":"","title":"","location":"","start":"YYYY-MM","end":"YYYY-MM or Present","context":"",
  "bullets":["..."]}],
 "skills":["..."], "education":[{"degree":"","institution":"","year":""}],
 "certifications":[{"name":"","issuer":"","status":"completed or in progress"}],
 "languages":["..."], "additional":["..."]}

Rules:
1. Copy every bullet EXACTLY as written in the resume, word for word. Do not shorten, merge, improve or summarise.
2. Use only dates written in the resume. If a month is missing use the year only ("2019"). If unknown, use "".
3. context: the one-line description under a job title (client, programme, team), copied exactly; else "".
4. skills: individual skills or tools as listed (split comma or bullet lists into separate items).
5. additional: projects, volunteering, hobbies or awards listed outside the jobs, copied exactly.
6. Never add anything that is not in the resume. Leave a field empty rather than guess.
7. Inside any text value, write a double quote as \\"."""


def propose(text: str, call: Callable[[list[dict]], str]) -> dict:
    """Ask the local model for a structured proposal of one resume."""
    from modules.local_analyzer import extract_json, recover_analysis  # noqa: F401  (same JSON tolerance)
    msgs = [{"role": "system", "content": EXTRACT_SYSTEM},
            {"role": "user", "content": f"RESUME:\n{text[:MAX_RESUME_CHARS]}\n\nReturn ONLY the JSON object."}]
    last = ""
    for attempt in range(3):
        raw = call(msgs) if attempt == 0 else call(msgs, temperature=0.3)
        try:
            data = extract_json(raw)
            if isinstance(data.get("roles"), list):
                return data
            last = "the reply has no 'roles' list"
        except ValueError as e:
            last = str(e)
        msgs += [{"role": "assistant", "content": (raw or "{}")[:3000]},
                 {"role": "user", "content": f"That was invalid ({last}). Return ONLY the corrected JSON object."}]
    raise RuntimeError(f"the model could not read this resume ({last})")


# --------------------------------------------------------------------------- #
# Verification and merging (pure Python)
# --------------------------------------------------------------------------- #

def _norm(s: str) -> str:
    s = (s or "").replace("\u2013", "-").replace("\u2014", "-").replace("\u2011", "-")
    return re.sub(r"\s+", " ", re.sub(r"^[\s\-•*·▪◦]+", "", s)).strip().lower()


def _words(s: str) -> list[str]:
    return re.findall(r"[a-z0-9+#&./%-]+", _norm(s))


def verify_bullet(bullet: str, source: str) -> str:
    """'verified' if the bullet is in the resume text (word for word, or 90%+ of its words in one passage),
    else 'reworded'."""
    b, src = _norm(bullet), _norm(source)
    if not b:
        return "reworded"
    if b in src:
        return "verified"
    bw = _words(bullet)
    sw = _words(source)
    if not bw:
        return "reworded"
    n = len(bw)
    window = set()
    best = 0.0
    for i in range(0, max(1, len(sw) - n + 1)):  # sliding window the size of the bullet (+20%)
        window = set(sw[i:i + int(n * 1.2) + 1])
        best = max(best, sum(1 for w in bw if w in window) / n)
        if best >= 0.9:
            return "verified"
    return "reworded"


def _year_ok(value: str, source: str) -> bool:
    m = re.search(r"(19|20)\d{2}", value or "")
    return (not value) or value.strip().lower() in ("present", "current") or (bool(m) and m.group(0) in source)


def _similar(a: str, b: str) -> bool:
    wa, wb = set(_words(a)), set(_words(b))
    return bool(wa and wb) and len(wa & wb) / len(wa | wb) >= 0.8


def merge(proposals: list[tuple[dict, str, str]]) -> dict:
    """Merge proposals from several resumes into one review list. Each item: (proposal, source_text, file name)."""
    roles: list[dict] = []
    skills: list[dict] = []
    education: list[dict] = []
    certs: list[dict] = []
    languages: list[str] = []
    additional: list[dict] = []
    for prop, src, fname in proposals:
        for r in prop.get("roles", []) or []:
            if not isinstance(r, dict) or not (r.get("employer") or r.get("title")):
                continue
            emp, start = _norm(r.get("employer", "")), _ym(r.get("start", "")) or None
            # same employer + same start date = same role, even if the title is written differently
            match = next((x for x in roles if _norm(x["employer"]) == emp and x["start_year"] == start and
                          (start or _similar(x["title"], r.get("title", "")))), None)
            if not match:
                match = {"employer": r.get("employer", "").strip(), "title": r.get("title", "").strip(),
                         "location": (r.get("location") or "").strip(), "start": (r.get("start") or "").strip(),
                         "end": (r.get("end") or "").strip(), "context": (r.get("context") or "").strip(),
                         "start_year": start, "bullets": [], "sources": [],
                         "dates_found": _year_ok(r.get("start", ""), src) and _year_ok(r.get("end", ""), src)}
                roles.append(match)
            match["sources"] = sorted(set(match["sources"]) | {fname})
            for b in r.get("bullets", []) or []:
                b = re.sub(r"^[\s\-•*·▪◦]+", "", str(b)).strip()
                if not b:
                    continue
                status = verify_bullet(b, src)
                twin = next((x for x in match["bullets"] if _similar(x["text"], b)), None)
                if twin:   # same fact in two resumes: keep the verified / longer wording
                    if (status, len(b)) > (twin["status"], len(twin["text"])) and status == "verified":
                        twin.update(text=b, status=status)
                    twin["sources"] = sorted(set(twin["sources"]) | {fname})
                else:
                    match["bullets"].append({"text": b, "status": status, "sources": [fname]})
        for s in prop.get("skills", []) or []:
            s = str(s).strip()
            if s and not any(_norm(s) == _norm(x["text"]) for x in skills):
                skills.append({"text": s, "found": _norm(s) in _norm(src)})
        for e in prop.get("education", []) or []:
            if isinstance(e, dict) and e.get("degree") and not any(_norm(e["degree"]) == _norm(x["degree"]) for x in education):
                education.append({"degree": e.get("degree", ""), "institution": e.get("institution", ""),
                                  "year": str(e.get("year", "") or "")})
        for c in prop.get("certifications", []) or []:
            if isinstance(c, dict) and c.get("name") and not any(_norm(c["name"]) == _norm(x["name"]) for x in certs):
                st = "in progress" if "progress" in str(c.get("status", "")).lower() else "completed"
                certs.append({"name": c["name"], "issuer": c.get("issuer", ""), "status": st})
        for lang in prop.get("languages", []) or []:
            if str(lang).strip() and _norm(lang) not in map(_norm, languages):
                languages.append(str(lang).strip())
        for a in prop.get("additional", []) or []:
            a = str(a).strip()
            if a and not any(_similar(a, x["text"]) for x in additional):
                additional.append({"text": a, "status": verify_bullet(a, src)})
    roles.sort(key=lambda r: (r["start"] or ""), reverse=True)
    return {"roles": roles, "skills": skills, "education": education, "certifications": certs,
            "languages": languages, "additional": additional}


# --------------------------------------------------------------------------- #
# Building the profile pack
# --------------------------------------------------------------------------- #

def _ym(v: str) -> str:
    v = (v or "").strip()
    if v.lower() in ("present", "current", "now", ""):
        return "Present" if v else ""
    m = re.match(r"^((19|20)\d{2})(?:[-/.](\d{1,2}))?", v)
    if m:
        return f"{m.group(1)}-{int(m.group(3)):02d}" if m.group(3) else m.group(1)
    months = {m: i for i, m in enumerate("jan feb mar apr may jun jul aug sep oct nov dec".split(), 1)}
    m2 = re.match(r"^([A-Za-z]{3})[a-z]*\.?\s+((19|20)\d{2})", v)
    if m2 and m2.group(1).lower() in months:
        return f"{m2.group(2)}-{months[m2.group(1).lower()]:02d}"
    return v


def build_pack(st: dict) -> tuple[dict, dict, str]:
    """From confirmed setup state -> (fact_bank, resume_profile, domain_toml_text)."""
    basics, conf = st["basics"], st["confirmed"]
    roles_fb, used_ids = [], set()
    for i, r in enumerate(conf["roles"]):
        rid = slugify(f"{r['employer']}_{(_ym(r['start']) or 'x')[:4]}") or f"role_{i}"
        while rid in used_ids:
            rid += "_b"
        used_ids.add(rid)
        bullets = [{"id": f"{rid[:12]}_{k + 1:02d}", "text": b, "category": ""} for k, b in enumerate(r["bullets"])]
        roles_fb.append({"role_id": rid, "employer": r["employer"], "title": r["title"], "location": r.get("location", ""),
                         "start": _ym(r["start"]), "end": _ym(r["end"]) or "Present", "track": "both",
                         "context": r.get("context", ""), "bullets": bullets})
    skills: dict[str, list[str]] = {}
    for s in conf["skills"]:
        skills.setdefault(s.get("section", "tools"), []).append(s["text"].strip() + LEVELS.get(s.get("level", "expert"), ""))
    edu = conf["education"][0] if conf["education"] else {}
    education = {"degree": edu.get("degree", ""), "institution": edu.get("institution", ""), "year": edu.get("year", "")}
    if conf.get("final_year_project"):
        education["final_year_project"] = {"description": conf["final_year_project"]}
    fb = {"schema_version": "2.0", "last_updated": datetime.now().strftime("%Y-%m-%d"),
          "candidate_meta": {"name": basics["name"].upper()}, "roles": roles_fb,
          "skills": {"core": {k: v for k, v in skills.items() if v}} if skills else {},
          "education": education,
          "certifications": [{"name": c["name"], "issuer": c.get("issuer", ""), "status": c["status"]} for c in conf["certifications"]],
          "additional": [{"id": f"add_{k + 1:03d}", "text": a, "category": "additional"} for k, a in enumerate(conf["additional"])],
          "languages": conf.get("languages", []), "removed_items": {"items": []}}

    sections = [f"core.{k}" for k in SECTIONS if skills.get(k)]
    targets = []
    for v in st["variants"]:
        include = [r for r in roles_fb if r["role_id"] in v["roles"]] or roles_fb
        sel = {}
        for n, r in enumerate(include):
            ids = [b["id"] for b in r["bullets"]]
            first, last = n == 0, n == len(include) - 1 and n > 0
            lo, hi = (3, 8) if first else ((1, 3) if last else (2, 5))
            sel[r["role_id"]] = {"lead_with": ids[:2], "then": ids[2:], "min_bullets": min(lo, len(ids)), "max_bullets": hi}
        targets.append({
            "target_id": slugify(v["label"]) or f"variant_{len(targets) + 1}", "label": v["label"],
            "matches_job_titles": [t.strip() for t in v.get("titles", "").split(",") if t.strip()],
            "headline_options": [v.get("headline") or v["label"]], "summary": v.get("summary", "").strip(),
            "role_order": [r["role_id"] for r in include], "bullet_selection": sel,
            "skills_sections_order": sections,
            "include_additional": [a["id"] for a in fb["additional"]],
            "include_education_project": bool(conf.get("final_year_project"))})
    cert_lines = [f"{c['name']}" + (f" — {c['issuer']}" if c.get("issuer") else "") +
                  (" — in progress" if c["status"] == "in progress" else "") for c in conf["certifications"]]
    rp = {"schema_version": "1.0",
          "candidate": {k: basics.get(k, "") for k in ("name", "phone", "email", "linkedin", "location")},
          "flags": {"include_languages": False, "google_pm_cert_on_all_resumes": True},
          "target_roles": targets, "shared_blocks": {"certifications": cert_lines}}
    rp["candidate"]["name"] = basics["name"].upper()
    pack = (PACKS_DIR / f"{st.get('pack', 'general')}.toml")
    domain = pack.read_text(encoding="utf-8") if pack.exists() else ""
    return fb, rp, domain


def write_pack(slug: str, st: dict, template_src: Optional[Path] = None) -> list[str]:
    """Validate, then write the profile files. Returns a list of problems (empty = success)."""
    import tomllib
    from modules.profile_pack import build_profile
    fb, rp, domain = build_pack(st)
    try:
        build_profile(slug, fb, rp, tomllib.loads(domain) if domain else {})
    except SystemExit as e:
        return [str(e)]
    if not any(r["bullets"] for r in fb["roles"]):
        return ["Keep at least one bullet in at least one role."]
    d = folder(slug)
    d.mkdir(parents=True, exist_ok=True)
    (d / "fact_bank.json").write_text(json.dumps(fb, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    (d / "resume_profile.json").write_text(json.dumps(rp, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    (d / "domain.toml").write_text(domain or "", encoding="utf-8")
    s = st.get("search", {})
    lines = ["# What the daily run searches for. Edited by the control centre's Settings page.", "",
             f"roles = {json.dumps(s.get('roles') or [v['label'] for v in st['variants']])}",
             f"locations = {json.dumps(s.get('locations', []))}", f"country = {json.dumps(s.get('country', ''))}",
             'sites = ["indeed", "linkedin"' + (', "naukri"]' if s.get("country", "").lower() == "india" else "]"),
             "scan_company_portals = true", 'title_match = "broad"', "max_age_days = 30", "full_refresh_every_days = 7",
             "", "[schedule]", "enabled = false", 'time = "06:30"',
             'days = ["mon", "tue", "wed", "thu", "fri", "sat", "sun"]', "", "[notify]", 'ntfy_topic = ""']
    (d / "search.toml").write_text("\n".join(lines) + "\n", encoding="utf-8")
    tpl = d / "resume_template.docx"
    if template_src and template_src.exists():
        from modules.resume_builder import make_template
        make_template(template_src, tpl)
    elif not tpl.exists():
        shutil.copy(EXAMPLE / "resume_template.docx", tpl)
    return []


def use_profile(slug: str) -> None:
    """Select this profile in config.local.toml (kept off GitHub)."""
    from modules.settings import ROOT
    p = ROOT / "config.local.toml"
    text = p.read_text(encoding="utf-8") if p.exists() else ""
    if re.search(r"^\s*profile\s*=", text, re.M):
        text = re.sub(r'^\s*profile\s*=.*$', f'profile = "{slug}"', text, flags=re.M)
    elif "[pipeline]" in text:
        text = text.replace("[pipeline]", f'[pipeline]\nprofile = "{slug}"', 1)
    else:
        text = text.rstrip() + f'\n\n[pipeline]\nprofile = "{slug}"\n'
    p.write_text(text.lstrip(), encoding="utf-8")
