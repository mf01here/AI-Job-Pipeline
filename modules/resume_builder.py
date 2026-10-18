#!/usr/bin/env python3
"""
Phase 5/6: Resume files (DOCX + PDF) and the application ledger
===============================================================

Input : data/tailored_drafts/*.json              (Phase 4: which bullets, in which order, final wording)
        profiles/<name>/                         (profile pack: header, skills, education, certificates)
        profiles/<name>/resume_template.docx     (your own resume's page setup and styles; body is empty)
Output: data/resumes/<company>-<title>-<hash8>/<Name>_<Company>_<Title>.docx  (+ .pdf)
        data/application_ledger.csv               (one row per job; your own edits are kept)

Honesty: every line printed comes from the Fact Bank, from resume_profile.json, or from the guarded Phase 4
draft. Nothing is written by a model here.

Layout: single column, standard headings, real text (no text boxes or images) - ATS friendly. Fonts, colours
and spacing match your MachineDesign resume.

Usage
    python modules/resume_builder.py --make-template "~/Desktop/My_Resume.docx"   # once: your layout -> template
    python modules/resume_builder.py                  # build resumes for every draft
    python modules/resume_builder.py --only Quest     # just matching companies/titles
    python modules/resume_builder.py --no-pdf         # Word files only
"""
from __future__ import annotations

import argparse
import csv
import json
import logging
import os
import re
import shutil
import subprocess
import sys
from datetime import datetime
from pathlib import Path
from typing import Optional

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from modules.profile_pack import Profile, load_profile  # noqa: E402
from modules.settings import DATA_DIR, profile_dir  # noqa: E402

log = logging.getLogger("resume")

DRAFTS_DIR = DATA_DIR / "tailored_drafts"
RESUMES_DIR = DATA_DIR / "resumes"
LEDGER = DATA_DIR / "application_ledger.csv"
LEDGER_COLUMNS = ["job_id", "company", "title", "location", "score", "variant", "status", "applied_on",
                  "notes", "resume_docx", "resume_pdf", "application_url", "generated_at", "content_hash"]

# Look of your MachineDesign resume
FONT = "Calibri"
NAVY, GREY, BLACK = "1F3557", "444444", "000000"
MONTHS = "Jan Feb Mar Apr May Jun Jul Aug Sep Oct Nov Dec".split()
ACRONYMS = {"cad", "plm", "sap", "kpi", "erp", "dfm", "dfa", "gd&t", "cae", "fea"}


# --------------------------------------------------------------------------- #
# Small formatting helpers (python-docx)
# --------------------------------------------------------------------------- #

def _docx():
    try:
        import docx  # noqa: F401
        return docx
    except ImportError:
        raise SystemExit("The python-docx package is missing. Install it once with:\n    pip install python-docx")


def _twips(n: int):
    from docx.shared import Twips
    return Twips(n)


def _run(par, text: str, size: float, bold=False, italic=False, color: Optional[str] = None, spacing: int = 0):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
    r = par.add_run(text)
    f = r.font
    f.name, f.size, f.bold, f.italic = FONT, Pt(size), bold, italic
    rpr = r._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(a), FONT)
    if color:
        f.color.rgb = RGBColor.from_string(color)
    if spacing:
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), str(spacing))
        rpr.append(sp)
    return r


def _par(doc, before=0, after=0, line: Optional[int] = None, center=False):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after = _twips(before), _twips(after)
    if line:
        pf.line_spacing = line / 240
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def _bottom_border(par, color=NAVY):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    ppr = par._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:color", color), ("w:sz", "6"), ("w:space", "3")):
        b.set(qn(k), v)
    bdr.append(b)
    ppr.insert(0, bdr)


def _bullet_num_id(doc) -> Optional[str]:
    """The numbering definition that draws bullets in the template (taken from your own resume)."""
    from docx.oxml.ns import qn
    try:
        numbering = doc.part.numbering_part.element
    except (KeyError, NotImplementedError, AttributeError):
        return None
    bullet_abstract = set()
    for ab in numbering.findall(qn("w:abstractNum")):
        lvl0 = ab.find(qn("w:lvl"))
        fmt = lvl0.find(qn("w:numFmt")) if lvl0 is not None else None
        if fmt is not None and fmt.get(qn("w:val")) == "bullet":
            bullet_abstract.add(ab.get(qn("w:abstractNumId")))
    for num in numbering.findall(qn("w:num")):
        ref = num.find(qn("w:abstractNumId"))
        if ref is not None and ref.get(qn("w:val")) in bullet_abstract:
            return num.get(qn("w:numId"))
    return None


def _bullet(doc, text: str, num_id: Optional[str]):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    p = _par(doc, after=24, line=220)
    if num_id:
        try:
            p.style = doc.styles["List Paragraph"]
        except KeyError:
            pass
        ppr = p._p.get_or_add_pPr()
        numpr = OxmlElement("w:numPr")
        ilvl, nid = OxmlElement("w:ilvl"), OxmlElement("w:numId")
        ilvl.set(qn("w:val"), "0")
        nid.set(qn("w:val"), num_id)
        numpr.append(ilvl)
        numpr.append(nid)
        ppr.insert(1 if ppr.find(qn("w:pStyle")) is not None else 0, numpr)
        p.paragraph_format.space_before, p.paragraph_format.space_after = _twips(0), _twips(24)
        _run(p, text, 9.5)
    else:  # no bullet definition in the template: a plain dash keeps it readable for ATS
        p.paragraph_format.left_indent = _twips(280)
        p.paragraph_format.first_line_indent = _twips(-180)
        _run(p, "- " + text, 9.5)
    return p


def _heading(doc, text: str):
    p = _par(doc, before=118, after=62)
    p.paragraph_format.keep_with_next = True
    _bottom_border(p)
    _run(p, text.upper(), 10.5, bold=True, color=NAVY, spacing=24)


def _label_line(doc, label: str, text: str):
    p = _par(doc, after=50, line=232)
    _run(p, f"{label}  ", 9.5, bold=True, color=NAVY)
    _run(p, text, 9.5)


def _month(ym: str) -> str:
    m = re.match(r"^(\d{4})-(\d{2})", str(ym or ""))
    if m:
        return f"{MONTHS[int(m.group(2)) - 1]} {m.group(1)}"
    return str(ym or "").strip() or ""


def _section_label(key: str) -> str:
    name = key.split(".", 1)[-1]
    words = name.replace("_and_", " & ").replace("_", " ").split()
    return " ".join(w.upper() if w.lower() in ACRONYMS else (w if w == "&" else w.capitalize()) for w in words)


# --------------------------------------------------------------------------- #
# Template
# --------------------------------------------------------------------------- #

def make_template(source: Path, dest: Path) -> None:
    """Copy your resume's page setup, styles and bullet definitions into an empty template."""
    docx = _docx()
    doc = docx.Document(str(source))
    body = doc.element.body
    for child in list(body):
        if not child.tag.endswith("}sectPr"):
            body.remove(child)
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dest))


# --------------------------------------------------------------------------- #
# Building one resume
# --------------------------------------------------------------------------- #

def summary_text(target_cfg: dict) -> str:
    """Profile paragraph assembled from resume_profile.json (your approved wording); never model-written."""
    if target_cfg.get("summary"):
        return target_cfg["summary"].strip()
    parts = target_cfg.get("summary_ingredients", [])
    picks = target_cfg.get("summary_from") or list(range(min(3, len(parts))))
    chosen = [parts[i].strip().rstrip(".") for i in picks if 0 <= i < len(parts)]
    if not chosen:
        return ""
    headline = (target_cfg.get("headline_options") or [""])[0]
    label = headline.split("|")[0].strip() or target_cfg.get("label", "").strip() or "Engineer"
    first = chosen[0]
    if re.match(r"^\d", first):          # "12+ years across ..." -> "Mechanical Design Engineer with 12+ years across ..."
        first = f"{label} with {first}"
    return ". ".join([first, *chosen[1:]]) + "."


def build_docx(draft: dict, profile: Profile, template: Path, out_path: Path) -> Path:
    docx = _docx()
    doc = docx.Document(str(template))
    num_id = _bullet_num_id(doc)
    rp = profile.resume_profile
    tid = (draft.get("target") or {}).get("id")
    target_cfg = next((t for t in rp.get("target_roles", []) if t.get("target_id") == tid), {})
    target = profile.targets.get(tid) or profile.default_target()
    cand = rp.get("candidate") or profile.fact_bank.get("candidate_meta", {})
    flags = rp.get("flags", {})

    # Header
    p = _par(doc, after=40, center=True)
    _run(p, cand.get("name", ""), 20, bold=True, color=NAVY, spacing=20)
    headline = (target_cfg.get("headline_options") or [""])[0]
    if headline:
        p = _par(doc, after=60, center=True)
        _run(p, headline, 10.5, bold=True, color=NAVY)
    contact = "  •  ".join(x for x in (cand.get("location"), cand.get("phone"), cand.get("email"),
                                       cand.get("linkedin")) if x)
    p = _par(doc, after=160, center=True)
    _run(p, contact, 9, color=GREY)

    # Profile
    summary = summary_text(target_cfg)
    if summary:
        _heading(doc, "Profile")
        p = _par(doc, after=60, line=250)
        _run(p, summary, 10)

    # Skills (the variant's skill sections, in its order)
    sections = target.skills_sections
    if sections:
        _heading(doc, "Technical Skills")
        for key in sections:
            items = [ln.text for ln in profile.lines if ln.kind == "skill" and ln.section == key]
            if items:
                _label_line(doc, _section_label(key), "  •  ".join(items))

    # Experience
    _heading(doc, "Experience")
    for sec in draft.get("sections", []):
        if sec.get("role_id") == "additional":
            continue
        role = profile.role(sec["role_id"])
        p = _par(doc, before=100, after=0)
        p.paragraph_format.keep_with_next = True
        from docx.enum.text import WD_TAB_ALIGNMENT
        p.paragraph_format.tab_stops.add_tab_stop(_twips(10800), WD_TAB_ALIGNMENT.RIGHT)
        _run(p, role.get("title", sec.get("title", "")), 10.5, bold=True, color=BLACK)
        where = ", ".join(x for x in (role.get("employer", ""), role.get("location", "")) if x)
        if where and not role.get("optional"):
            _run(p, f"  |  {where}", 10.5, color=BLACK)
        _run(p, "\t", 10.5)
        dates = f"{_month(role.get('start'))} – {_month(role.get('end')) or 'Present'}"
        _run(p, dates, 9.5, bold=True, color=GREY)
        if role.get("context") and not role.get("optional"):
            p = _par(doc, before=20, after=40)
            p.paragraph_format.keep_with_next = True
            _run(p, role["context"], 9.5, italic=True, color=GREY)
        bullets = sec.get("bullets", [])
        grouped = any(b.get("project") for b in bullets)
        if grouped:   # keep each project's bullets together, projects in order of first appearance
            order = list(dict.fromkeys(b.get("project", "") for b in bullets))
            bullets = [b for proj in order for b in bullets if b.get("project", "") == proj]
        current = None
        for b in bullets:
            if grouped and b.get("project") and b["project"] != current:
                current = b["project"]
                p = _par(doc, before=66, after=24)
                p.paragraph_format.keep_with_next = True
                _run(p, current, 9.5, bold=True, color=NAVY)
            _bullet(doc, b["final"], num_id)

    # Education, certificates, languages, additional
    _heading(doc, "Education & Additional")
    edu = profile.fact_bank.get("education", {})
    if edu.get("degree"):
        _label_line(doc, "Education", ", ".join(str(x) for x in (edu["degree"], edu.get("institution"), edu.get("year")) if x))
    proj = edu.get("final_year_project") or {}
    if target_cfg.get("include_education_project") and proj.get("description"):
        _label_line(doc, "Final-year project", proj["description"])
    certs = (rp.get("shared_blocks") or {}).get("certifications") or []
    if not flags.get("google_pm_cert_on_all_resumes", True):
        certs = [c for c in certs if "in progress" not in c.lower()]
    if certs:
        _label_line(doc, "Certifications", "  •  ".join(certs))
    if flags.get("include_languages") and profile.fact_bank.get("languages"):
        _label_line(doc, "Languages", "  •  ".join(profile.fact_bank["languages"]))
    extra = [b["final"] for sec in draft.get("sections", []) if sec.get("role_id") == "additional"
             for b in sec.get("bullets", [])]
    if extra:
        _label_line(doc, "Additional", "  •  ".join(extra))

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #

def _soffice() -> Optional[str]:
    for c in (shutil.which("soffice"), shutil.which("libreoffice"),
              "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if c and Path(c).exists():
            return c
    return None


def to_pdf(docx_path: Path) -> Optional[Path]:
    """DOCX -> PDF with LibreOffice if installed, else Microsoft Word (via docx2pdf). None if neither works."""
    pdf = docx_path.with_suffix(".pdf")
    office = _soffice()
    if office:
        try:
            subprocess.run([office, "--headless", "--convert-to", "pdf", "--outdir", str(docx_path.parent),
                            str(docx_path)], check=True, capture_output=True, timeout=180)
            if pdf.exists():
                return pdf
        except (subprocess.SubprocessError, OSError) as e:
            log.debug("LibreOffice conversion failed: %s", e)
    try:
        from docx2pdf import convert  # uses Microsoft Word on a Mac
        convert(str(docx_path), str(pdf))
        if pdf.exists():
            return pdf
    except Exception as e:  # noqa: BLE001
        log.debug("docx2pdf conversion failed: %s", e)
    return None


# --------------------------------------------------------------------------- #
# Ledger
# --------------------------------------------------------------------------- #

def read_ledger() -> dict[str, dict]:
    if not LEDGER.exists():
        return {}
    with open(LEDGER, newline="", encoding="utf-8") as f:
        return {row["job_id"]: row for row in csv.DictReader(f) if row.get("job_id")}


def write_ledger(rows: dict[str, dict]) -> None:
    LEDGER.parent.mkdir(parents=True, exist_ok=True)
    tmp = LEDGER.with_suffix(".csv.tmp")
    with open(tmp, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=LEDGER_COLUMNS, extrasaction="ignore")
        w.writeheader()
        order = {"applied": 1, "interview": 0, "offer": 0}
        for row in sorted(rows.values(), key=lambda r: (order.get(r.get("status", ""), 2),
                                                         -float(r.get("score") or 0))):
            w.writerow(row)
    os.replace(tmp, LEDGER)


def update_ledger(rows: dict[str, dict], draft: dict, docx_path: Path, pdf_path: Optional[Path]) -> None:
    job = draft["job"]
    old = rows.get(job["job_id"], {})
    keep_status = old.get("status") and old.get("status") != "resume_ready"   # never overwrite your own progress
    rows[job["job_id"]] = {
        **old,
        "job_id": job["job_id"], "company": job["company"], "title": job["title"], "location": job.get("location", ""),
        "score": f"{draft.get('score', 0):.1f}", "variant": (draft.get("target") or {}).get("label", ""),
        "status": old["status"] if keep_status else "resume_ready",
        "applied_on": old.get("applied_on", ""), "notes": old.get("notes", ""),
        "resume_docx": str(docx_path.relative_to(ROOT)) if docx_path.is_relative_to(ROOT) else str(docx_path),
        "resume_pdf": (str(pdf_path.relative_to(ROOT)) if pdf_path and pdf_path.is_relative_to(ROOT)
                       else (str(pdf_path) if pdf_path else "")),
        "application_url": job.get("application_url", ""),
        "generated_at": datetime.now().isoformat(timespec="seconds"), "content_hash": job.get("content_hash", ""),
    }


# --------------------------------------------------------------------------- #
# Main
# --------------------------------------------------------------------------- #

def _safe(s: str, n: int = 40) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", s).strip("_")[:n]


def main() -> int:
    ap = argparse.ArgumentParser(description="Phase 5: resume files and application ledger")
    ap.add_argument("--profile", help="profile folder under profiles/ (default: config.toml)")
    ap.add_argument("--make-template", metavar="DOCX", help="turn one of your resumes into the layout template, then exit")
    ap.add_argument("--only", action="append", metavar="TEXT", help="only drafts whose company/title contains TEXT")
    ap.add_argument("--limit", type=int)
    ap.add_argument("--job-id", action="append", metavar="ID", help="only this job (repeatable)")
    ap.add_argument("--no-pdf", action="store_true", help="Word files only")
    ap.add_argument("--new-only", action="store_true", help="skip jobs whose resume is newer than their draft")
    ap.add_argument("-v", "--verbose", action="store_true")
    args = ap.parse_args()
    logging.basicConfig(level=logging.DEBUG if args.verbose else logging.INFO, format="%(message)s", stream=sys.stderr)

    folder = profile_dir(args.profile)
    template = folder / "resume_template.docx"
    if args.make_template:
        src = Path(os.path.expanduser(args.make_template))
        if not src.exists():
            raise SystemExit(f"{src} not found.")
        make_template(src, template)
        log.info("Template saved -> %s\n(Only the page setup, fonts and bullet style were kept; the text was removed.)",
                 template)
        return 0
    if not template.exists():
        raise SystemExit(f"No resume template yet. Create it once from your favourite resume:\n"
                         f'    python modules/resume_builder.py --make-template "path/to/your_resume.docx"')

    profile = load_profile(folder)
    drafts = sorted(DRAFTS_DIR.glob("*.json"))
    picked = []
    for path in drafts:
        d = json.loads(path.read_text(encoding="utf-8"))
        if not d.get("sections"):
            log.info("Skipping %s (old draft format; redraft it with: python modules/cloud_generator.py --fresh)", path.name)
            continue
        if d.get("fact_bank_hash") and d["fact_bank_hash"] != profile.fact_hash:
            log.warning("Note: %s was drafted from an older Fact Bank; consider redrafting.", path.name)
        if args.job_id and d["job"]["job_id"] not in set(args.job_id):
            continue
        hay = f"{d['job']['company']} {d['job']['title']}".lower()
        if args.only and not all(t.lower() in hay for t in args.only):
            continue
        picked.append((path, d))
    if args.limit is not None:
        picked = picked[:args.limit]
    if not picked:
        if args.new_only:
            log.info("No new drafts: every resume is up to date.")
            return 0
        raise SystemExit(f"No drafts to build in {DRAFTS_DIR}. Run Phase 4 first: python modules/cloud_generator.py")

    name = _safe((profile.resume_profile.get("candidate") or {}).get("name", "Resume").title(), 30)
    rows = read_ledger()
    pdf_ok = pdf_missing = built = 0
    for path, d in picked:
        job = d["job"]
        out = RESUMES_DIR / path.stem / f"{name}_{_safe(job['company'], 25)}_{_safe(job['title'])}.docx"
        if args.new_only and out.exists() and out.stat().st_mtime >= path.stat().st_mtime:
            continue
        build_docx(d, profile, template, out)
        pdf = None if args.no_pdf else to_pdf(out)
        if pdf:
            pdf_ok += 1
        elif not args.no_pdf:
            pdf_missing += 1
        update_ledger(rows, d, out, pdf)
        built += 1
        log.info("  %-28s %-45s -> %s%s", job["company"][:28], job["title"][:45], out.name, "  (+pdf)" if pdf else "")
    write_ledger(rows)
    log.info("\n%d resume(s) saved in %s%s", built, RESUMES_DIR,
             f" ({len(picked) - built} already up to date)" if len(picked) > built else "")
    if pdf_missing:
        log.info("PDFs were not made for %d resume(s): install LibreOffice (free) or run  pip install docx2pdf  "
                 "(uses Microsoft Word), then run this again.", pdf_missing)
    log.info("Application ledger -> %s", LEDGER)
    return 0


if __name__ == "__main__":
    sys.exit(main())
